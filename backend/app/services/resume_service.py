"""Resume ingestion and analysis.

Deterministic by design: the resume text is extracted from the uploaded file,
skills are detected by matching a known vocabulary against that text, and the
scores are computed from real structural signals. No model invents skills, and a
mention is never treated as mastery.

On success the detected skills are written to `skills` + `skill_evidence`, which
is what makes the Phase 06 skill-gap engine operate on real evidence.
"""

from __future__ import annotations

import hashlib
import io
import logging
import re
from datetime import datetime, timezone
from typing import Any

from app.core.config import Settings
from app.core.errors import ApiError
from app.integrations.storage import SupabaseStorageClient
from app.repositories.resumes import ResumesRepository
from app.repositories.skills import SkillsRepository
from app.schemas.resume import (
    ResumeAnalysis,
    ResumeImprovements,
    ResumeListResponse,
    ResumeOut,
    ResumeSection,
    SectionStatus,
)
from app.services.skill_gap_service import normalize_skill_key

logger = logging.getLogger("careeros.resume")

MAX_UPLOAD_BYTES = 10 * 1024 * 1024
RESUME_BUCKET = "resumes"
_SUPPORTED_SUFFIXES = (".pdf", ".docx")

_ACTION_VERBS = (
    "led", "built", "designed", "implemented", "improved", "reduced", "increased",
    "developed", "managed", "created", "optimized", "delivered", "launched",
    "automated", "deployed", "migrated", "architected", "analysed", "analyzed",
    "shipped", "maintained", "collaborated", "mentored",
)

_SECTION_KEYWORDS = {
    "experience": ("experience", "employment", "work history", "internship"),
    "education": ("education", "b.tech", "bachelor", "master", "university", "college", "degree"),
    "skills": ("skills", "technologies", "tech stack", "technical skills"),
    "projects": ("projects", "project"),
    "summary": ("summary", "objective", "profile"),
}

# Alternate spellings so common resume phrasing still matches the vocabulary.
_SKILL_ALIASES: dict[str, tuple[str, ...]] = {
    "machine-learning": ("machine learning", "ml", "supervised learning"),
    "deep-learning": ("deep learning", "neural network", "neural networks"),
    "mlops": ("mlops", "ml ops", "model deployment"),
    "python": ("python",),
    "sql": ("sql", "mysql", "postgresql", "postgres"),
    "javascript": ("javascript", "js", "es6"),
    "typescript": ("typescript", "ts"),
    "react": ("react", "react.js", "reactjs"),
    "html": ("html", "html5"),
    "css": ("css", "css3", "tailwind"),
    "kubernetes": ("kubernetes", "k8s"),
    "docker": ("docker", "containers"),
    "excel": ("excel", "spreadsheets", "pivot tables"),
    "data-visualization": ("data visualization", "tableau", "power bi", "matplotlib"),
    "statistics": ("statistics", "statistical", "regression", "hypothesis testing"),
    "a-b-testing": ("a/b testing", "ab testing", "split testing"),
    "rest-apis": ("rest", "rest api", "restful", "api"),
    "git": ("git", "github", "gitlab", "version control"),
    "system-design": ("system design", "distributed systems", "architecture"),
    "accessibility": ("accessibility", "a11y", "wcag", "aria"),
    "pandas": ("pandas",),
    "pytorch": ("pytorch", "torch"),
}


class ResumeService:
    def __init__(
        self,
        resumes: ResumesRepository,
        skills: SkillsRepository,
        storage: SupabaseStorageClient,
        settings: Settings,
    ) -> None:
        self._resumes = resumes
        self._skills = skills
        self._storage = storage
        self._settings = settings

    async def upload(
        self,
        claims: dict[str, Any],
        access_token: str,
        *,
        filename: str,
        content_type: str | None,
        content: bytes,
    ) -> ResumeOut:
        _validate_upload(filename, content)

        text = extract_text(filename, content_type, content)
        if not text.strip():
            raise ApiError(
                422,
                "no_text_extracted",
                "No readable text was found in that file. Scanned images are not supported.",
            )

        user_id = str(claims["sub"])
        content_hash = hashlib.sha256(content).hexdigest()

        storage_path = await self._storage.upload(
            bucket=RESUME_BUCKET,
            path=f"{user_id}/{content_hash}{_suffix(filename)}",
            content=content,
            content_type=content_type or "application/octet-stream",
            access_token=access_token,
        )

        analysis = await self._analyse(claims, user_id, text)

        row = await self._resumes.create(
            claims,
            user_id,
            filename=filename,
            mime_type=content_type,
            size_bytes=len(content),
            content_hash=content_hash,
            storage_path=storage_path,
            parse_status="parsed",
            analysis_status="complete",
            extracted_text=text,
            extracted_data={
                "word_count": analysis.word_count,
                "detected_skills": analysis.detected_skills,
                "target_role": analysis.target_role,
                "storage_stored": storage_path is not None,
            },
            analysis=analysis.model_dump(mode="json"),
            ats_score=analysis.ats_score,
        )

        # Close the loop: evidence recorded now feeds the skill-gap engine.
        await self._record_skill_evidence(claims, user_id, row["id"], analysis.detected_skills)

        return _resume_out(row)

    async def latest(self, claims: dict[str, Any], user_id: str) -> ResumeOut | None:
        row = await self._resumes.get_latest(claims, user_id)
        return _resume_out(row) if row else None

    async def list_resumes(self, claims: dict[str, Any], user_id: str) -> ResumeListResponse:
        rows = await self._resumes.list_for_user(claims, user_id)
        return ResumeListResponse(resumes=[_resume_out(row) for row in rows])

    async def delete(self, claims: dict[str, Any], user_id: str, resume_id: str) -> None:
        deleted = await self._resumes.delete(claims, user_id, resume_id)
        if not deleted:
            raise ApiError(404, "resume_not_found", "That resume does not exist.")

    # ── analysis ──────────────────────────────────────────────────────────

    async def _analyse(
        self, claims: dict[str, Any], user_id: str, text: str
    ) -> ResumeAnalysis:
        now = datetime.now(timezone.utc)
        lowered = text.lower()

        vocabulary = await self._skills.list_skill_vocabulary(claims)
        by_key = {str(row["skill_key"]): str(row["skill_name"]) for row in vocabulary}
        detected_keys = _detect_skills(lowered, by_key)

        profile = await self._skills.get_profile_context(claims, user_id)
        target_role = (profile or {}).get("target_role_name") or None

        required_names: list[str] = []
        missing: list[str] = []
        role_fit: int | None = None
        if target_role:
            requirement = await self._skills.get_latest_role_requirement(claims, target_role)
            if requirement is not None:
                required = await self._skills.list_required_skills(
                    claims, str(requirement["id"])
                )
                required_names = [str(item["skill_name"]) for item in required]
                required_keys = [str(item["skill_key"]) for item in required]
                if required_keys:
                    matched_count = sum(1 for key in required_keys if key in detected_keys)
                    role_fit = round(matched_count / len(required_keys) * 100)
                    missing = [
                        by_key.get(key, key)
                        for key in required_keys
                        if key not in detected_keys
                    ]

        detected_names = [by_key[key] for key in sorted(detected_keys) if key in by_key]

        signals = _signals(text, lowered)
        ats_score = _ats_score(signals, role_fit, len(detected_keys))
        quality_score = _quality_score(signals)
        sections = _sections(signals, ats_score, quality_score, role_fit)
        improvements = _improvements(signals, detected_names, required_names, missing, target_role)

        return ResumeAnalysis(
            ats_score=ats_score,
            quality_score=quality_score,
            role_fit_score=role_fit,
            sections=sections,
            detected_skills=detected_names,
            missing_skills=missing,
            improvements=improvements,
            word_count=signals["word_count"],
            target_role=target_role,
            analyzed_at=now,
        )

    async def _record_skill_evidence(
        self,
        claims: dict[str, Any],
        user_id: str,
        resume_id: str,
        detected_names: list[str],
    ) -> None:
        for name in detected_names:
            key = normalize_skill_key(name)
            skill_id = await self._skills.upsert_skill(
                claims,
                user_id,
                skill_key=key,
                skill_name=name,
                evidence_summary="Found in uploaded resume.",
            )
            await self._skills.add_evidence(
                claims,
                user_id,
                skill_id=skill_id or None,
                skill_key=key,
                skill_name=name,
                source_type="resume",
                source_ref=resume_id,
                detail="Skill name detected in resume text.",
            )


# ── parsing ───────────────────────────────────────────────────────────────


def _suffix(filename: str) -> str:
    lowered = filename.lower()
    return ".pdf" if lowered.endswith(".pdf") else ".docx"


def _validate_upload(filename: str, content: bytes) -> None:
    lowered = filename.lower()
    if not lowered.endswith(_SUPPORTED_SUFFIXES):
        raise ApiError(
            415,
            "unsupported_file",
            "Only PDF and DOCX resumes can be analysed.",
        )
    if not content:
        raise ApiError(400, "empty_file", "The uploaded file is empty.")
    if len(content) > MAX_UPLOAD_BYTES:
        raise ApiError(413, "file_too_large", "Resumes must be 10MB or smaller.")


def extract_text(filename: str, content_type: str | None, content: bytes) -> str:
    """Extract plain text from a PDF or DOCX file."""
    lowered = filename.lower()
    if lowered.endswith(".pdf") or content_type == "application/pdf":
        return _extract_pdf(content)
    return _extract_docx(content)


def _extract_pdf(content: bytes) -> str:
    try:
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(content))
        return "\n".join((page.extract_text() or "") for page in reader.pages)
    except ApiError:
        raise
    except Exception as exc:  # malformed/corrupt PDF
        logger.warning("resume_pdf_parse_failed error=%s", type(exc).__name__)
        raise ApiError(422, "unreadable_file", "That PDF could not be read.") from exc


def _extract_docx(content: bytes) -> str:
    try:
        from docx import Document

        document = Document(io.BytesIO(content))
        parts = [paragraph.text for paragraph in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                parts.extend(cell.text for cell in row.cells)
        return "\n".join(parts)
    except ApiError:
        raise
    except Exception as exc:
        logger.warning("resume_docx_parse_failed error=%s", type(exc).__name__)
        raise ApiError(422, "unreadable_file", "That DOCX could not be read.") from exc


# ── deterministic analysis helpers ────────────────────────────────────────


def _detect_skills(lowered_text: str, vocabulary: dict[str, str]) -> set[str]:
    detected: set[str] = set()
    for key, name in vocabulary.items():
        candidates = (name.lower(), *_SKILL_ALIASES.get(key, ()))
        for candidate in candidates:
            if _contains_term(lowered_text, candidate):
                detected.add(key)
                break
    return detected


def _contains_term(haystack: str, term: str) -> bool:
    """Word-boundary match that tolerates punctuation around the term."""
    pattern = r"(?<![a-z0-9])" + re.escape(term) + r"(?![a-z0-9])"
    return re.search(pattern, haystack) is not None


def _signals(text: str, lowered: str) -> dict[str, Any]:
    words = re.findall(r"[A-Za-z][A-Za-z0-9+#./-]*", text)
    bullets = [line for line in text.splitlines() if line.strip()[:1] in {"-", "•", "*", "·"}]
    numbers = re.findall(r"(?<![a-z0-9])\d[\d,.]*\s*(?:%|percent|k\b|lakh|cr\b|million)?", lowered)
    quantified = len(re.findall(r"\d+\s*%|\d+\s*(?:x|times)\b|₹\s?\d|\$\s?\d", lowered))
    verb_hits = sum(1 for word in words if word.lower() in _ACTION_VERBS)
    return {
        "raw": text,
        "lowered": lowered,
        "word_count": len(words),
        "has_email": re.search(r"[\w.+-]+@[\w-]+\.[\w.-]+", text) is not None,
        "has_phone": re.search(r"(?:\+?\d[\d\s().-]{7,}\d)", text) is not None,
        "has_link": re.search(r"(linkedin\.com|github\.com|https?://)", lowered) is not None,
        "sections_present": {
            name: any(_contains_term(lowered, kw) for kw in keywords)
            for name, keywords in _SECTION_KEYWORDS.items()
        },
        "bullets": len(bullets),
        "quantified": quantified,
        "numbers": len(numbers),
        "verb_hits": verb_hits,
    }


def _ats_score(signals: dict[str, Any], role_fit: int | None, detected_count: int) -> int:
    score = 0
    if signals["has_email"]:
        score += 12
    if signals["has_phone"]:
        score += 10
    if signals["has_link"]:
        score += 8
    score += min(20, sum(1 for present in signals["sections_present"].values() if present) * 5)

    word_count = signals["word_count"]
    if 300 <= word_count <= 900:
        score += 15
    elif word_count >= 150:
        score += 7

    if role_fit is not None:
        score += round(role_fit * 0.35)
    else:
        score += min(35, detected_count * 5)

    return max(0, min(100, score))


def _quality_score(signals: dict[str, Any]) -> int:
    score = 0
    if signals["quantified"] >= 2:
        score += 30
    elif signals["quantified"] == 1:
        score += 15
    elif signals["numbers"] >= 3:
        score += 10

    if signals["verb_hits"] >= 8:
        score += 25
    elif signals["verb_hits"] >= 4:
        score += 15
    elif signals["verb_hits"] >= 1:
        score += 7

    if signals["bullets"] >= 8:
        score += 20
    elif signals["bullets"] >= 3:
        score += 12

    if 300 <= signals["word_count"] <= 900:
        score += 25
    elif signals["word_count"] >= 150:
        score += 12

    return max(0, min(100, score))


def _status(score: int) -> SectionStatus:
    if score >= 80:
        return "good"
    if score >= 60:
        return "warn"
    return "bad"


def _sections(
    signals: dict[str, Any], ats_score: int, quality_score: int, role_fit: int | None
) -> list[ResumeSection]:
    contact = 0
    contact += 40 if signals["has_email"] else 0
    contact += 30 if signals["has_phone"] else 0
    contact += 30 if signals["has_link"] else 0

    present = signals["sections_present"]
    keyword_score = min(100, sum(1 for value in present.values() if value) * 20)
    experience_score = min(100, signals["verb_hits"] * 8 + signals["quantified"] * 10)
    education_score = 100 if present.get("education") and present.get("projects") else 55

    sections = [
        ResumeSection(label="Formatting & Contact", score=contact, status=_status(contact)),
        ResumeSection(label="Skills & Keywords", score=keyword_score, status=_status(keyword_score)),
        ResumeSection(
            label="Experience Descriptions",
            score=experience_score,
            status=_status(experience_score),
        ),
        ResumeSection(
            label="Education & Projects", score=education_score, status=_status(education_score)
        ),
        ResumeSection(label="Action Verbs & Metrics", score=quality_score, status=_status(quality_score)),
    ]
    if role_fit is not None:
        sections.append(ResumeSection(label="Role Fit", score=role_fit, status=_status(role_fit)))
    return sections


def _improvements(
    signals: dict[str, Any],
    detected: list[str],
    required: list[str],
    missing: list[str],
    target_role: str | None,
) -> ResumeImprovements:
    critical: list[str] = []
    recommended: list[str] = []
    optional: list[str] = []

    if missing:
        preview = ", ".join(f"'{name}'" for name in missing[:3])
        critical.append(
            f"Add {preview} — required for {target_role or 'your target role'} and not found in your resume."
        )
    if signals["quantified"] < 2:
        critical.append(
            "Quantify your impact with numbers (for example 'reduced inference time by 30%')."
        )

    if not signals["sections_present"].get("summary"):
        recommended.append("Add a 2–3 sentence summary aligned to your target role.")
    if not signals["has_link"]:
        recommended.append("Add a LinkedIn or GitHub link so reviewers can verify your work.")
    if required and len(detected) < len(required):
        recommended.append("Mirror the exact skill names used in the target job description.")

    if signals["word_count"] < 300:
        optional.append("Your resume is short — expand the most relevant projects.")
    if signals["word_count"] > 900:
        optional.append("Consider trimming to the most role-relevant content.")
    if signals["bullets"] < 3:
        optional.append("Use bullet points for achievements rather than paragraphs.")

    return ResumeImprovements(critical=critical, recommended=recommended, optional=optional)


def _resume_out(row: dict[str, Any]) -> ResumeOut:
    stored = row.get("analysis")
    analysis = ResumeAnalysis.model_validate(stored) if stored else None
    return ResumeOut(
        id=str(row["id"]),
        filename=str(row["filename"]),
        mime_type=row.get("mime_type"),
        size_bytes=row.get("size_bytes"),
        version=int(row["version"]),
        parse_status=str(row["parse_status"]),  # type: ignore[arg-type]
        analysis_status=str(row["analysis_status"]),  # type: ignore[arg-type]
        storage_path=row.get("storage_path"),
        content_hash=str(row["content_hash"]),
        target_role=(analysis.target_role if analysis else None),
        analysis=analysis,
        created_at=row.get("created_at"),
    )
